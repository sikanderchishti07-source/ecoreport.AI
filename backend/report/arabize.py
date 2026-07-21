# -*- coding: utf-8 -*-
"""Builds master_template_ar.docx from the English master template.

Strategy: the English template is the single source of truth for structure
(tables, Jinja tags, fields, images). This module opens a freshly built EN
template, replaces every translatable run using the exact-match table in
i18n_ar.AR, then applies right-to-left formatting:

- w:bidi on every body/header paragraph (RTL paragraph direction)
- explicit LEFT alignment flipped to RIGHT (centre/justify preserved)
- w:bidiVisual on every table (columns flow right-to-left)
- complex-script font set to Amiri on all runs (+ szCs mirror of sz)
- w:rtl on runs that contain Arabic characters

Jinja tags, chemical symbols, method IDs and numeric placeholders are left
untouched; the Word bidi algorithm displays embedded Latin/digits correctly
inside RTL paragraphs.
"""
import os
import re
import tempfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn

from .i18n_ar import AR

HERE = os.path.dirname(os.path.abspath(__file__))
OUT_AR = os.path.join(HERE, "master_template_ar.docx")

ARABIC_RE = re.compile(r"[\u0600-\u06FF]")
AR_FONT = "Amiri"


def _set_pPr_bidi(paragraph):
    pPr = paragraph._p.get_or_add_pPr()
    if pPr.find(qn("w:bidi")) is None:
        bidi = pPr.makeelement(qn("w:bidi"), {})
        pPr.insert(0, bidi)


def _set_table_bidi(table):
    tblPr = table._tbl.tblPr
    if tblPr.find(qn("w:bidiVisual")) is None:
        bv = tblPr.makeelement(qn("w:bidiVisual"), {})
        tblPr.insert(0, bv)


def _style_run_ar(run):
    rPr = run._r.get_or_add_rPr()
    fonts = rPr.find(qn("w:rFonts"))
    if fonts is None:
        fonts = rPr.makeelement(qn("w:rFonts"), {})
        rPr.insert(0, fonts)
    fonts.set(qn("w:cs"), AR_FONT)
    sz = rPr.find(qn("w:sz"))
    if sz is not None and rPr.find(qn("w:szCs")) is None:
        szCs = rPr.makeelement(qn("w:szCs"), {})
        szCs.set(qn("w:val"), sz.get(qn("w:val")))
        rPr.append(szCs)
    if ARABIC_RE.search(run.text or ""):
        if rPr.find(qn("w:rtl")) is None:
            rPr.append(rPr.makeelement(qn("w:rtl"), {}))


def _process_paragraph(p):
    for run in p.runs:
        t = run.text
        if t and t in AR:
            run.text = AR[t]
        _style_run_ar(run)
    _set_pPr_bidi(p)
    if p.alignment == WD_ALIGN_PARAGRAPH.LEFT:
        p.alignment = WD_ALIGN_PARAGRAPH.RIGHT


def _process_container(container):
    for p in container.paragraphs:
        _process_paragraph(p)
    for t in getattr(container, "tables", []):
        _set_table_bidi(t)
        for row in t.rows:
            for cell in row.cells:
                _process_container(cell)


def build_ar(out_path: str = OUT_AR) -> str:
    """Build the Arabic master template. Rebuilds the EN template into a temp
    file first so the transform always starts from the current EN source."""
    from .template_builder import build as build_en
    with tempfile.TemporaryDirectory() as td:
        tmp_en = os.path.join(td, "en.docx")
        build_en(tmp_en)
        doc = Document(tmp_en)

        _process_container(doc)
        for section in doc.sections:
            for part in (section.header, section.footer,
                         section.first_page_header, section.first_page_footer):
                if part is not None:
                    _process_container(part)

        # Complex-script default for the whole document
        for style_name in ["Normal", "Heading 1", "Heading 2", "Heading 3",
                           "List Bullet", "List Bullet 2"]:
            try:
                st = doc.styles[style_name]
                rPr = st.element.get_or_add_rPr()
                fonts = rPr.find(qn("w:rFonts"))
                if fonts is None:
                    fonts = rPr.makeelement(qn("w:rFonts"), {})
                    rPr.insert(0, fonts)
                fonts.set(qn("w:cs"), AR_FONT)
            except KeyError:
                pass

        doc.save(out_path)
    return out_path


if __name__ == "__main__":
    print("Arabic template written:", build_ar())
